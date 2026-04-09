from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ページ余白設定
section = doc.sections[0]
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(2)
section.right_margin = Cm(2)

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_section_title(doc, text, bg_hex='2E4057'):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    # 背景色はWordの段落シェーディングで
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), bg_hex)
    pPr.append(shd)

def add_table(doc, headers, rows, header_bg='48A999'):
    col_count = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=col_count)
    table.style = 'Table Grid'
    table.autofit = False

    # ヘッダー行
    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(10)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_bg(cell, header_bg)

    # データ行
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = str(val)
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)
            if c_idx == 0:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if r_idx % 2 == 1:
                set_cell_bg(cell, 'F0F8F6')

    return table

# ===== タイトル =====
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title.add_run('出演・出店者 報告一覧')
tr.bold = True
tr.font.size = Pt(16)
title.paragraph_format.space_after = Pt(12)

# ===== 1. ステージ出演 =====
add_section_title(doc, '■ ステージ出演', '2E4057')
stage_rows = [
    ['①', '天然デンネンズ', '歌・ギター'],
    ['②', 'POTATO SOUND（ポテトサウンド）', 'バンド'],
    ['③', 'ドラマーせん', 'ドラムソロ'],
    ['④', 'BLACK haedman（ブラックヘッドマン）', 'バンド'],
    ['⑤', 'DST', 'ダンス'],
]
t1 = add_table(doc, ['No.', '出演者名', 'ジャンル'], stage_rows, '2E4057')
# 列幅
for i, w in enumerate([Cm(1.5), Cm(8), Cm(5)]):
    for cell in t1.columns[i].cells:
        cell.width = w

# ===== 2. 手づくり市（ハンドメイド・ワークショップ・加工品） =====
add_section_title(doc, '■ 手づくり市　ハンドメイド・ワークショップ・加工品その他', '3D6B84')
craft_rows = [
    ['①',  'MISORASIDO★こども手づくり市', '', ''],
    ['②',  'Bouquet de colza（菜の花の花束）', '', ''],
    ['③',  'わはは牧場 shop＆cafe がぶう', '', ''],
    ['④',  '天然ツムクム', '', ''],
    ['⑤',  'ゆゆたまサロン', '', ''],
    ['⑥',  '工房HIROCRAFT', '', ''],
    ['⑦',  'neutraldsign', '', ''],
    ['⑧',  'KOMUTENこむてん', '', ''],
    ['⑨',  'cacoiropeace', '', ''],
    ['⑩',  'flat.a.flat', '', ''],
    ['⑪',  'Rikk.', '', ''],
    ['⑫',  "WoodworkingassistanceM's maechi.thb", '', ''],
    ['⑬',  '菓子工房ルーエプラス', '', ''],
    ['⑭',  'エマーノ＆ふくみ', '', ''],
    ['⑮',  'fukupuku', '※フィーカさんの隣', ''],
    ['⑯',  'ping pong ハニー', '', '電源'],
    ['⑰',  '正法寺ノミノイチ', '', '電源・火気'],
    ['⑱',  'かめの湯', '', '電源'],
]
t2 = add_table(doc, ['No.', '店舗名', '備考', '設備'], craft_rows, '3D6B84')
for i, w in enumerate([Cm(1.5), Cm(8), Cm(3.5), Cm(2.5)]):
    for cell in t2.columns[i].cells:
        cell.width = w

# ===== 3. FOOD & ドリンク =====
add_section_title(doc, '■ 手づくり市　FOOD＆ドリンク', '5E8C61')
food_rows = [
    ['⑲', '食力屋Tarzan', '', '火気'],
    ['⑳', 'Cafe Fika', '', '火気・電源'],
    ['㉑', 'アルトス・ヴィレッジ', '', '火気'],
]
t3 = add_table(doc, ['No.', '店舗名', '備考', '設備'], food_rows, '5E8C61')
for i, w in enumerate([Cm(1.5), Cm(8), Cm(3.5), Cm(2.5)]):
    for cell in t3.columns[i].cells:
        cell.width = w

# ===== 4. キッチンカー =====
add_section_title(doc, '■ 手づくり市　キッチンカー', '7B6B9A')
kc_rows = [
    ['㉒', 'チョコバナナ',          '軽・幅４M', 'スィーツ',           '火気'],
    ['㉓', '翁ゆきちゃん号',        '軽・幅４M', 'たこ焼き',           '火気'],
    ['㉔', 'こここめキッチン',      '軽・幅４M', 'スィーツ',           '火気・電源'],
    ['㉕', 'Ueda Base',             '大・幅５M', 'トルティーヤ・ポテト', '火気'],
    ['㉖', 'DanRanfoodservice',     '大・幅５M', 'ピザ他',             '火気・電源'],
    ['㉗', '唐揚げ専門 はっぴ商店', '大・幅５M', 'からあげ',           '火気・電源'],
    ['㉘', 'ゴリラの台所',          '大・幅５M', '焼きうどん・唐揚',   '火気・電源'],
    ['㉙', '屋台ニシカラ',          '大・幅５M', 'からあげ',           '火気・電源'],
]
t4 = add_table(doc, ['No.', '店舗名', '車種・幅', '販売品目', '設備'], kc_rows, '7B6B9A')
for i, w in enumerate([Cm(1.5), Cm(5.5), Cm(2.5), Cm(4), Cm(2.5)]):
    for cell in t4.columns[i].cells:
        cell.width = w

out_path = '/home/user/biglabo-simulation/出演出店者一覧.docx'
doc.save(out_path)
print(f'Saved: {out_path}')
